"""상담데이터 등록용 대화내용 엑셀 읽기 / 양식 만들기.

엑셀 형식 (첫 번째 시트)
- 1행 머리글: 화자 | 대화내용   (머리글이 없으면 A열=화자, B열=대화내용으로 봄)
- 화자: 고객 / 상담사 (손님·상담원·CUSTOMER·AGENT 도 인식)
- 빈 행은 건너뜀. 순서는 행 순서 그대로
"""
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

MAX_LINES = 200          # 등록 API(ConsultDataIn.utterances) 한도와 같음
MAX_TEXT = 1000          # 등록 API(LineIn.text) 한도와 같음
MAX_ERRORS = 5

SPEAKERS = {
    "CUSTOMER": {"고객", "손님", "customer", "cust", "c"},
    "AGENT": {"상담사", "상담원", "직원", "agent", "a"},
}
SPEAKER_HEADERS = {"화자", "구분", "발화자", "speaker"}
TEXT_HEADERS = {"대화내용", "대화", "내용", "발화", "text"}


class DialogExcelError(ValueError):
    pass


def _cell(v) -> str:
    return "" if v is None else str(v).strip()


def _speaker(v: str) -> str | None:
    key = v.replace(" ", "").lower()
    for code, names in SPEAKERS.items():
        if key in names:
            return code
    return None


def _find_columns(row: list[str]) -> tuple[int, int] | None:
    """머리글 행이면 (화자 열, 대화 열) 위치 반환."""
    norm = [c.replace(" ", "").lower() for c in row]
    sp = next((i for i, c in enumerate(norm) if c in SPEAKER_HEADERS), None)
    tx = next((i for i, c in enumerate(norm) if c in TEXT_HEADERS), None)
    return (sp, tx) if sp is not None and tx is not None else None


def parse(data: bytes) -> dict:
    """엑셀 바이트 → {"lines": [{"speaker", "text"}], "sheet": 시트명}. 형식 오류는 DialogExcelError."""
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError):
        raise DialogExcelError("엑셀 파일(.xlsx)을 읽을 수 없습니다. 엑셀에서 '.xlsx' 형식으로 저장한 파일을 올려 주세요.") from None
    try:
        ws = wb.worksheets[0]
        rows = [(i, [_cell(v) for v in r]) for i, r in enumerate(ws.iter_rows(values_only=True), start=1)]
    finally:
        wb.close()

    rows = [(i, r) for i, r in rows if any(r)]
    if not rows:
        raise DialogExcelError("엑셀 첫 번째 시트에 대화내용이 없습니다.")
    cols = _find_columns(rows[0][1])
    if cols:
        rows = rows[1:]
    sp_col, tx_col = cols or (0, 1)

    lines, errors = [], []
    for i, r in rows:
        sp_raw = r[sp_col] if sp_col < len(r) else ""
        text = r[tx_col] if tx_col < len(r) else ""
        if not sp_raw and not text:
            continue
        speaker = _speaker(sp_raw)
        if not speaker:
            errors.append(f"{i}행: 화자 '{sp_raw or '(빈칸)'}' 은(는) '고객' 또는 '상담사' 여야 합니다.")
        elif not text:
            errors.append(f"{i}행: 대화내용이 비어 있습니다.")
        elif len(text) > MAX_TEXT:
            errors.append(f"{i}행: 대화내용이 {MAX_TEXT}자를 넘습니다. ({len(text)}자)")
        else:
            lines.append({"speaker": speaker, "text": text})
        if len(errors) >= MAX_ERRORS:
            break
    if errors:
        raise DialogExcelError(" / ".join(errors))
    if not lines:
        raise DialogExcelError("엑셀에서 읽은 대화가 없습니다.")
    if len(lines) > MAX_LINES:
        raise DialogExcelError(f"대화는 최대 {MAX_LINES}건까지 등록할 수 있습니다. (엑셀 {len(lines)}건)")
    return {"lines": lines, "sheet": ws.title}


def template() -> bytes:
    """대화내용 입력 양식 (.xlsx): 머리글 + 예시 3행 + 화자 선택 목록."""
    wb = Workbook()
    ws = wb.active
    ws.title = "대화내용"
    ws.append(["화자", "대화내용"])
    for row in [("상담사", "안녕하세요, OO은행 고객센터 상담사입니다. 무엇을 도와드릴까요?"),
                ("고객", "모바일 앱에서 로그인이 계속 실패해요."),
                ("상담사", "불편을 드려 죄송합니다. 어떤 오류 메시지가 표시되나요?")]:
        ws.append(row)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF9")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 90
    for r in ws.iter_rows(min_row=2, max_col=2):
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"고객,상담사"', allow_blank=True,
                        errorTitle="화자", error="고객 또는 상담사를 선택하세요.")
    ws.add_data_validation(dv)
    dv.add(f"A2:A{MAX_LINES + 1}")
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
